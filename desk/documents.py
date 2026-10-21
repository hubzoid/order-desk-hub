"""Turn a message and its attachments into numbered source elements.

Every later step points at these elements: an extracted line cites `E12`,
and the case view shows exactly that line of the PDF, that row of the sheet
or that line of the email. Sources:

* the message text: one element per non-empty line;
* CSV and plain text files: one element per row or line;
* spreadsheets (XLSX): one element per non-empty row, with its cells and
  the sheet and cell range (openpyxl, read only, formulas as stored values);
* PDFs and images: Docling, in its own process with a deadline, one document
  at a time (text items, and each table row with its cells). OCR only for
  images, or a PDF without a text layer.

Files are checked first: size, page count, encryption, macros and type.
Extracted text is untrusted data, never instructions.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import subprocess
import sys
import tempfile
import threading
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

SUPPORTED = {
    "application/pdf": "pdf", "image/png": "image", "image/jpeg": "image", "text/csv": "csv", "text/plain": "text",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}
_DOCLING_LOCK = threading.Lock()     # one Docling document at a time in this process


class ParseError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


@dataclass
class Limits:
    max_file_bytes: int = 20 * 1024 * 1024
    max_pages: int = 100
    timeout_seconds: int = 180


@dataclass
class Parsed:
    label: str
    parser: str
    elements: list[dict] = field(default_factory=list)
    pages: int | None = None


def media_kind(filename: str, media_type: str | None) -> str | None:
    mt = (media_type or "").split(";")[0].strip().lower()
    if mt in SUPPORTED:
        return SUPPORTED[mt]
    ext = Path(filename or "").suffix.lower()
    return {".pdf": "pdf", ".png": "image", ".jpg": "image", ".jpeg": "image", ".csv": "csv", ".txt": "text",
            ".xlsx": "xlsx"}.get(ext)


def message_elements(text: str, label: str = "message") -> Parsed:
    out = []
    for n, line in enumerate((text or "").splitlines(), 1):
        if line.strip():
            out.append({"kind": "text", "loc": f"{label} line {n}", "line": n, "text": line.strip()})
    return Parsed(label=label, parser="text", elements=out)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        d = Decimal(repr(v))
        return format(d.normalize(), "f") if d == d.to_integral_value() or len(repr(v)) < 18 else repr(v)
    if isinstance(v, Decimal):
        return format(v, "f")
    return str(v).strip()


def _xlsx(data: bytes, label: str) -> Parsed:
    if b"vbaProject.bin" in data or b"xl/vbaProject" in data[:65536]:
        raise ParseError("file_unsupported", f"{label} contains macros. Save it without macros (XLSX).")
    try:
        zipfile.ZipFile(io.BytesIO(data)).testzip()
    except zipfile.BadZipFile:
        raise ParseError("parse_failed", f"{label} is not a readable spreadsheet (it may be encrypted or damaged).") \
            from None
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise ParseError("parse_failed", f"{label} is not a readable spreadsheet.") from None
    out = []
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            if len(out) > 5000:
                raise ParseError("file_too_large", f"{label} has more than 5000 rows.")
            last = get_column_letter(len(cells))
            out.append({"kind": "row", "loc": f"{label} {ws.title}!A{r}:{last}{r}", "sheet": ws.title, "row": r,
                        "cells": cells, "text": " | ".join(c for c in cells)})
    wb.close()
    return Parsed(label=label, parser="openpyxl", elements=out)


def _csv(data: bytes, label: str) -> Parsed:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ParseError("parse_failed", f"{label} is not UTF-8 text.") from None
    out = []
    for r, row in enumerate(csv.reader(io.StringIO(text)), 1):
        cells = [c.strip() for c in row]
        if any(cells):
            out.append({"kind": "row", "loc": f"{label} row {r}", "row": r, "cells": cells, "text": " | ".join(cells)})
    return Parsed(label=label, parser="csv", elements=out)


def _pdf_precheck(data: bytes, label: str, limits: Limits) -> tuple[int, bool]:
    """(pages, has a text layer). Encrypted and over-long PDFs are refused
    before Docling starts."""
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseError("file_unsupported", f"{label} is encrypted. Send an unencrypted copy.")
        pages = len(reader.pages)
        text = "".join((reader.pages[i].extract_text() or "") for i in range(min(pages, 3)))
    except ParseError:
        raise
    except (PdfReadError, Exception):  # noqa: BLE001 - a damaged file
        raise ParseError("parse_failed", f"{label} is damaged or not a PDF.") from None
    if pages > limits.max_pages:
        raise ParseError("file_too_large", f"{label} has {pages} pages. The limit is {limits.max_pages}.")
    return pages, bool(text.strip())


def _docling(data: bytes, label: str, *, ocr: bool, limits: Limits, suffix: str) -> Parsed:
    with _DOCLING_LOCK, tempfile.TemporaryDirectory(prefix="desk-parse-") as tmp:
        path = Path(tmp) / f"document{suffix}"
        path.write_bytes(data)
        cmd = [sys.executable, "-m", "desk.docling_worker", str(path), "--max-pages", str(limits.max_pages)]
        if ocr:
            cmd.append("--ocr")
        env = dict(os.environ, PYTHONPATH=os.pathsep.join([str(Path(__file__).resolve().parents[1]),
                                                           os.environ.get("PYTHONPATH", "")]))
        try:
            done = subprocess.run(cmd, capture_output=True, text=True, timeout=limits.timeout_seconds, env=env,
                                  cwd=tmp, start_new_session=True)
        except subprocess.TimeoutExpired:
            raise ParseError("parse_timeout", f"{label} took longer than {limits.timeout_seconds} seconds to read. "
                             "Send a smaller document or one page at a time.") from None
        if done.returncode != 0:
            tail = (done.stderr or "").strip().splitlines()[-1:] or [""]
            if "ModuleNotFoundError" in done.stderr and "docling" in done.stderr:
                raise ParseError("parser_unavailable", "Docling is not installed. Install the parsing extra "
                                 "(pip install -e '.[parsing]' in the hub folder) and run `hubzoid-desk parser "
                                 "prefetch`.")
            if "LocalEntryNotFoundError" in done.stderr or "offline" in done.stderr.lower():
                raise ParseError("parser_unavailable", "Docling's model assets are not on this machine. Run "
                                 "`hubzoid-desk parser prefetch` once.")
            raise ParseError("parse_failed", f"{label} could not be read ({tail[0][:120]}).")
        try:
            out = json.loads(done.stdout.strip().splitlines()[-1])
        except (ValueError, IndexError):
            raise ParseError("parse_failed", f"{label} could not be read.") from None
    elements = []
    tables: dict[int, int] = {}
    for e in out["elements"]:
        page = e.get("page") or 1
        if e["kind"] == "row":
            if e.get("row") == 0:
                tables[page] = tables.get(page, 0) + 1
            t = tables.get(page, 1)
            elements.append({"kind": "row", "loc": f"{label} page {page} table {t} row {e['row'] + 1}", "page": page,
                             "row": e["row"] + 1, "cells": e["cells"], "text": e["text"]})
        else:
            n = sum(1 for x in elements if x.get("page") == page and x["kind"] == "text") + 1
            elements.append({"kind": "text", "loc": f"{label} page {page} line {n}" if suffix == ".pdf"
                             else f"{label} line {n}", "page": page, "text": e["text"]})
    return Parsed(label=label, parser=out.get("parser", "docling"), elements=elements, pages=out.get("pages"))


def parse_file(data: bytes, filename: str, media_type: str | None, *, limits: Limits | None = None) -> Parsed:
    limits = limits or Limits()
    label = Path(filename or "attachment").name
    if len(data) > limits.max_file_bytes:
        raise ParseError("file_too_large", f"{label} is larger than {limits.max_file_bytes} bytes.")
    kind = media_kind(filename, media_type)
    if kind is None:
        raise ParseError("file_unsupported", f"{label} is a {media_type or 'file'} this desk does not read. Send a PDF, "
                         "an image (PNG or JPEG), a spreadsheet (XLSX or CSV) or the order in the message text.")
    if kind == "text":
        try:
            return message_elements(data.decode("utf-8"), label=label)
        except UnicodeDecodeError:
            raise ParseError("parse_failed", f"{label} is not UTF-8 text.") from None
    if kind == "csv":
        return _csv(data, label)
    if kind == "xlsx":
        return _xlsx(data, label)
    if kind == "pdf":
        if not data.startswith(b"%PDF-"):
            raise ParseError("parse_failed", f"{label} is not a PDF.")
        pages, text_layer = _pdf_precheck(data, label, limits)
        parsed = _docling(data, label, ocr=not text_layer, limits=limits, suffix=".pdf")
        parsed.pages = parsed.pages or pages
        return parsed
    return _docling(data, label, ocr=True, limits=limits, suffix=Path(label).suffix.lower() or ".png")


def number(documents: list[Parsed]) -> list[dict]:
    """One flat, numbered element list across all documents: E1, E2, ..."""
    out, n = [], 0
    for d, doc in enumerate(documents, 1):
        for e in doc.elements:
            n += 1
            out.append(dict(e, id=f"E{n}", doc=d, doc_label=doc.label))
    return out


def sha256(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()
